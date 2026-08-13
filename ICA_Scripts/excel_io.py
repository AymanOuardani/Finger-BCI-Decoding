"""
excel_io.py

Thin openpyxl helper for writing analysis results as .xlsx workbooks into
Ressources/.

Kept deliberately small: one function writes one sheet from a header plus a list
of rows, replacing that sheet if it already exists and leaving every other sheet
of the workbook untouched. That is what lets several scripts (and several runs)
accumulate into the same workbook without clobbering each other.

Unlike the .ods tracking files, these workbooks hold no formulas, so they need
none of the LibreOffice prefix repair described in CLAUDE.md.
"""

import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def write_sheet(path, sheet_name, header, rows, freeze_header=True,
                number_format=None, column_width=None):
    """Write `rows` under `header` into `sheet_name` of the workbook at `path`.

    The workbook is created if missing. An existing sheet of the same name is
    replaced; all other sheets are preserved. Sheet names are truncated to the
    31-character Excel limit.

    number_format  optional {header_name: excel_format} applied to that column
    column_width   optional {header_name: width}; otherwise widths are derived
                   from the content
    """
    os.makedirs(os.path.dirname(path), exist_ok=True)
    sheet_name = str(sheet_name)[:31]

    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    if sheet_name in wb.sheetnames:
        wb.remove(wb[sheet_name])
    ws = wb.create_sheet(sheet_name)

    ws.append(list(header))
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append(list(row))

    number_format = number_format or {}
    column_width = column_width or {}
    for col_idx, name in enumerate(header, start=1):
        letter = get_column_letter(col_idx)
        if name in number_format:
            for cell in ws[letter][1:]:
                cell.number_format = number_format[name]
        if name in column_width:
            width = column_width[name]
        else:
            longest = max([len(str(name))] +
                          [len(str(r[col_idx - 1])) for r in rows
                           if col_idx <= len(r)] or [0])
            width = min(max(longest + 2, 9), 42)
        ws.column_dimensions[letter].width = width

    if freeze_header:
        ws.freeze_panes = "A2"

    _order_sheets(wb)
    wb.save(path)
    return path


def _order_sheets(wb):
    """Keep sheets alphabetically ordered so a growing workbook stays navigable."""
    for position, name in enumerate(sorted(wb.sheetnames)):
        wb.move_sheet(name, offset=position - wb.sheetnames.index(name))


def sheet_names(path):
    """Sheets currently in the workbook, or [] if it does not exist yet."""
    if not os.path.exists(path):
        return []
    return load_workbook(path, read_only=True).sheetnames
